"""Download stable Novani wholesale invoices and account ledger as Excel."""
from io import BytesIO

from django.contrib.auth.decorators import login_required
from django.http import HttpResponse
from django.shortcuts import get_object_or_404

from .dateutils import format_jalali
from .models import (
    WholesaleCustomer, WholesaleInvoice, WholesaleReceipt,
)
from .wholesale_logic_v119 import account


def _excel_text(value):
    """Prevent buyer-supplied names/notes from becoming executable Excel formulas."""
    value = str(value or "")
    return "'" + value if value.startswith(("=", "+", "-", "@")) else value


@login_required
def wholesale_export(request):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    customer_raw = request.GET.get("customer")
    selected = None
    if customer_raw:
        selected = get_object_or_404(WholesaleCustomer, pk=customer_raw)

    book = Workbook()
    invoice_tab = book.active
    invoice_tab.title = "Invoices"
    invoice_tab.append([
        "Invoice ID", "Jalali date", "Buyer", "Pieces",
        "Invoice amount toman", "Tailor account offset toman", "Note",
    ])
    item_tab = book.create_sheet("Items")
    item_tab.append([
        "Invoice ID", "Buyer", "Model / Color", "Size",
        "Pieces", "Unit price toman", "Line total toman",
    ])

    invoices = WholesaleInvoice.objects.select_related("customer").prefetch_related(
        "lines__color", "lines__size"
    ).order_by("date", "id")
    if selected:
        invoices = invoices.filter(customer=selected)
    for invoice in invoices:
        invoice_tab.append([
            invoice.pk, format_jalali(invoice.date), _excel_text(invoice.customer.name),
            invoice.quantity_total, invoice.total_amount,
            invoice.tailor_deducted_amount, _excel_text(invoice.note),
        ])
        for item in invoice.lines.all():
            item_tab.append([
                invoice.pk, _excel_text(invoice.customer.name), _excel_text(item.color.name), _excel_text(item.size.name),
                item.quantity, item.unit_price, item.amount,
            ])

    balance_tab = book.create_sheet("Buyer balances")
    balance_tab.append([
        "Buyer", "Gross invoiced", "Tailor offset", "Cash received",
        "Remaining", "Phone",
    ])
    customers = WholesaleCustomer.objects.filter(pk=selected.pk) if selected else WholesaleCustomer.objects.all()
    for buyer in customers:
        ledger = account(buyer)
        balance_tab.append([
            _excel_text(buyer.name), ledger["gross"], ledger["offset"], ledger["received"],
            ledger["balance"], _excel_text(buyer.phone),
        ])

    receipt_tab = book.create_sheet("Receipts")
    receipt_tab.append(["Jalali date", "Buyer", "Amount toman", "Note"])
    receipts = WholesaleReceipt.objects.select_related("customer").order_by("date", "pk")
    if selected:
        receipts = receipts.filter(customer=selected)
    for receipt in receipts:
        receipt_tab.append([
            format_jalali(receipt.date), _excel_text(receipt.customer.name),
            receipt.amount, _excel_text(receipt.note),
        ])

    for tab in book.worksheets:
        tab.freeze_panes = "A2"
        tab.auto_filter.ref = tab.dimensions
        for cell in tab[1]:
            cell.font = Font(color="FFFFFF", bold=True)
            cell.fill = PatternFill("solid", fgColor="243C5A")
            cell.alignment = Alignment(horizontal="center")
        for column in tab.columns:
            letter = column[0].column_letter
            tab.column_dimensions[letter].width = min(
                44, max(17, max(len(str(cell.value or "")) for cell in column) + 3)
            )
    output = BytesIO()
    book.save(output)
    response = HttpResponse(
        output.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = 'attachment; filename="novani-wholesale.xlsx"'
    return response
