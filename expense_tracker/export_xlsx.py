from collections import defaultdict
from io import BytesIO

import jdatetime
from django.contrib.auth.decorators import login_required
from django.db.models import Count, Sum
from django.http import HttpResponse
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from core.dateutils import format_jalali

from .models import DailyExpense, ReceivableEntry, ReceivablePerson
from .services import mellat_balance


HEADER_FILL = PatternFill("solid", fgColor="173B34")
HEADER_FONT = Font(color="FFFFFF", bold=True)
MONEY_FORMAT = '#,##0'


def _safe_text(value):
    text = str(value or "")
    if text[:1] in {"=", "+", "-", "@"}:
        return "'" + text
    return text


def _style_sheet(ws, *, money_columns=()):
    ws.sheet_view.rightToLeft = True
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(horizontal="right", vertical="center")
    for index in money_columns:
        for cell in ws[get_column_letter(index)][1:]:
            cell.number_format = MONEY_FORMAT


def _fit_columns(ws, widths=None):
    widths = widths or {}
    for idx in range(1, ws.max_column + 1):
        letter = get_column_letter(idx)
        if idx in widths:
            ws.column_dimensions[letter].width = widths[idx]
            continue
        max_len = 0
        for cell in ws[letter]:
            value = "" if cell.value is None else str(cell.value)
            max_len = max(max_len, len(value))
        ws.column_dimensions[letter].width = min(max(max_len + 2, 11), 36)


def _jalali_month_key(gregorian_date):
    j = jdatetime.date.fromgregorian(date=gregorian_date)
    return j.year, j.month


def _jalali_month_days(year, month):
    start = jdatetime.date(year, month, 1)
    if month == 12:
        nxt = jdatetime.date(year + 1, 1, 1)
    else:
        nxt = jdatetime.date(year, month + 1, 1)
    return (nxt.togregorian() - start.togregorian()).days


def _receivable_snapshot():
    total_outstanding = 0
    open_people = 0
    people_rows = []
    for person in ReceivablePerson.objects.all().order_by("name", "id"):
        totals = {
            row["kind"]: int(row["total"] or 0)
            for row in person.entries.values("kind").annotate(total=Sum("amount"))
        }
        claim = totals.get(ReceivableEntry.CLAIM, 0)
        payment = totals.get(ReceivableEntry.PAYMENT, 0)
        outstanding = claim - payment
        if outstanding > 0:
            total_outstanding += outstanding
            open_people += 1
        people_rows.append((person, claim, payment, outstanding))
    return total_outstanding, open_people, people_rows


def build_financial_workbook(*, today=None):
    today = today or timezone.localdate()
    today_j = jdatetime.date.fromgregorian(date=today)

    expenses = list(
        DailyExpense.objects.select_related("category").order_by("date", "created_at", "id")
    )
    receivable_entries = list(
        ReceivableEntry.objects.select_related("person").order_by("date", "created_at", "id")
    )

    wb = Workbook()

    # 1) Raw expense ledger
    ws = wb.active
    ws.title = "هزینه‌ها"
    ws.append([
        "ردیف",
        "تاریخ شمسی",
        "تاریخ میلادی",
        "دسته",
        "عنوان",
        "مبلغ (تومان)",
        "توضیح",
        "زمان ثبت",
    ])
    for idx, expense in enumerate(expenses, 1):
        ws.append([
            idx,
            format_jalali(expense.date),
            expense.date.isoformat(),
            _safe_text(expense.category.name),
            _safe_text(expense.title or expense.category.name),
            int(expense.amount or 0),
            _safe_text(expense.note),
            timezone.localtime(expense.created_at).strftime("%Y-%m-%d %H:%M:%S"),
        ])
    _style_sheet(ws, money_columns=(6,))
    _fit_columns(ws, {2: 14, 3: 14, 4: 18, 5: 28, 6: 18, 7: 36, 8: 21})

    # 2) Receivable / repayment ledger
    ws = wb.create_sheet("گردش طلب‌ها")
    ws.append([
        "ردیف",
        "تاریخ شمسی",
        "تاریخ میلادی",
        "شخص",
        "نوع",
        "مبلغ (تومان)",
        "اثر روی ملت اعمال شده",
        "توضیح",
        "زمان ثبت",
    ])
    for idx, entry in enumerate(receivable_entries, 1):
        ws.append([
            idx,
            format_jalali(entry.date),
            entry.date.isoformat(),
            _safe_text(entry.person.name),
            "طلب" if entry.kind == ReceivableEntry.CLAIM else "تسویه",
            int(entry.amount or 0),
            "بله" if entry.mellat_applied else "خیر (قدیمی)",
            _safe_text(entry.note),
            timezone.localtime(entry.created_at).strftime("%Y-%m-%d %H:%M:%S"),
        ])
    _style_sheet(ws, money_columns=(6,))
    _fit_columns(ws, {2: 14, 3: 14, 4: 22, 5: 12, 6: 18, 7: 23, 8: 36, 9: 21})

    # 3) Monthly expense summary (Jalali months)
    monthly = defaultdict(lambda: {"total": 0, "count": 0})
    for expense in expenses:
        key = _jalali_month_key(expense.date)
        monthly[key]["total"] += int(expense.amount or 0)
        monthly[key]["count"] += 1

    ws = wb.create_sheet("خلاصه ماهانه")
    ws.append(["ماه شمسی", "جمع خرج (تومان)", "تعداد تراکنش", "روزهای مبنای میانگین", "میانگین روزانه (تومان)"])
    for year, month in sorted(monthly):
        values = monthly[(year, month)]
        if year == today_j.year and month == today_j.month:
            elapsed_days = max(1, today_j.day)
        else:
            elapsed_days = _jalali_month_days(year, month)
        ws.append([
            f"{year}/{month:02d}",
            values["total"],
            values["count"],
            elapsed_days,
            int(round(values["total"] / elapsed_days)) if elapsed_days else 0,
        ])
    _style_sheet(ws, money_columns=(2, 5))
    _fit_columns(ws, {1: 14, 2: 20, 3: 16, 4: 22, 5: 24})

    # 4) Category summary, all recorded expenses
    ws = wb.create_sheet("خلاصه دسته‌ها")
    ws.append(["دسته", "تعداد تراکنش", "جمع خرج (تومان)", "سهم از کل خرج"])
    category_rows = list(
        DailyExpense.objects.values("category__name")
        .annotate(total=Sum("amount"), count=Count("id"))
        .order_by("-total", "category__name")
    )
    grand_total = sum(int(row["total"] or 0) for row in category_rows)
    for row in category_rows:
        total = int(row["total"] or 0)
        ws.append([
            _safe_text(row["category__name"]),
            int(row["count"] or 0),
            total,
            round(total * 100 / grand_total, 1) if grand_total else 0,
        ])
    _style_sheet(ws, money_columns=(3,))
    for cell in ws["D"][1:]:
        cell.number_format = '0.0"%"'
    _fit_columns(ws, {1: 22, 2: 17, 3: 20, 4: 17})

    # 5) Current financial snapshot
    total_outstanding, open_people, _people_rows = _receivable_snapshot()
    current_month_total = sum(
        int(expense.amount or 0)
        for expense in expenses
        if _jalali_month_key(expense.date) == (today_j.year, today_j.month)
    )
    today_total = sum(int(expense.amount or 0) for expense in expenses if expense.date == today)
    daily_average = int(round(current_month_total / max(1, today_j.day)))

    ws = wb.create_sheet("وضعیت فعلی")
    ws.append(["شاخص", "مقدار", "توضیح"])
    snapshot_rows = [
        ("تاریخ تهیه خروجی", format_jalali(today), ""),
        ("موجودی ملت", int(mellat_balance()), "مانده فعلی حساب ملت"),
        ("کل طلب باز", int(total_outstanding), "جمع طلب باقی‌مانده از همه اشخاص"),
        ("تعداد اشخاص دارای طلب باز", int(open_people), ""),
        ("جمع خرج ماه جاری", int(current_month_total), f"ماه {today_j.year}/{today_j.month:02d}"),
        ("میانگین خرج روزانه ماه", int(daily_average), f"تقسیم بر {today_j.day} روز گذشته ماه"),
        ("خرج امروز", int(today_total), ""),
        ("جمع کل خرج ثبت‌شده", int(sum(int(x.amount or 0) for x in expenses)), ""),
        ("تعداد کل تراکنش‌های خرج", len(expenses), ""),
        ("تعداد گردش‌های طلب/تسویه", len(receivable_entries), ""),
    ]
    for row in snapshot_rows:
        ws.append([_safe_text(row[0]), row[1], _safe_text(row[2])])
    _style_sheet(ws, money_columns=(2,))
    # Count/date rows should stay readable as ordinary values.
    for row_index in (2, 5, 10, 11):
        if row_index <= ws.max_row:
            ws.cell(row=row_index, column=2).number_format = "General"
    _fit_columns(ws, {1: 28, 2: 22, 3: 40})

    return wb


@login_required
def financial_export_xlsx(request):
    wb = build_financial_workbook()
    output = BytesIO()
    wb.save(output)
    output.seek(0)

    today = timezone.localdate()
    filename = f"kharj-financial-export-{today.isoformat()}.xlsx"
    response = HttpResponse(
        output.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    response["Cache-Control"] = "no-store"
    return response
