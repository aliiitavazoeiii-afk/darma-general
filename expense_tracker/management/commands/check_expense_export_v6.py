from io import BytesIO

from django.core.management.base import BaseCommand, CommandError
from django.test import RequestFactory
from django.template.loader import get_template
from django.urls import resolve
from openpyxl import load_workbook

from core.payment_source_v63 import SOURCE_MELAT, source_balance
from expense_tracker.models import DailyExpense, ReceivableEntry, ReceivablePerson


EXPECTED_SHEETS = [
    "هزینه‌ها",
    "گردش طلب‌ها",
    "خلاصه ماهانه",
    "خلاصه دسته‌ها",
    "وضعیت فعلی",
]


class Command(BaseCommand):
    help = "Read-only regression for the financial XLSX export V6."

    def _snapshot(self):
        return {
            "mellat": int(source_balance(SOURCE_MELAT) or 0),
            "expenses": DailyExpense.objects.count(),
            "receivable_entries": ReceivableEntry.objects.count(),
            "receivable_people": ReceivablePerson.objects.count(),
        }

    def handle(self, *args, **options):
        reports_source = get_template("expense_tracker/reports.html").template.source
        if "financial_export_xlsx" not in reports_source or "خروجی اکسل مالی" not in reports_source:
            raise CommandError("reports page is missing the financial Excel download action")

        before = self._snapshot()

        match = resolve("/reports/export.xlsx")
        request = RequestFactory().get("/reports/export.xlsx")
        request.user = type("AuthUser", (), {"is_authenticated": True})()
        response = match.func(request)

        if response.status_code != 200:
            raise CommandError(f"Excel export returned HTTP {response.status_code}")
        if response.get("Content-Type") != "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet":
            raise CommandError("Excel export content type is wrong")
        disposition = response.get("Content-Disposition", "")
        if "attachment" not in disposition or ".xlsx" not in disposition:
            raise CommandError("Excel export is missing an XLSX attachment disposition")
        if len(response.content) < 1000:
            raise CommandError("Excel export payload is unexpectedly small")

        try:
            workbook = load_workbook(BytesIO(response.content), read_only=True, data_only=False)
        except Exception as exc:
            raise CommandError(f"Excel export could not be opened by openpyxl: {exc}") from exc

        if workbook.sheetnames != EXPECTED_SHEETS:
            raise CommandError(
                f"Excel sheet set/order mismatch: {workbook.sheetnames!r}"
            )

        if workbook["هزینه‌ها"].max_row != DailyExpense.objects.count() + 1:
            raise CommandError("expense sheet row count does not match the database")
        if workbook["گردش طلب‌ها"].max_row != ReceivableEntry.objects.count() + 1:
            raise CommandError("receivable sheet row count does not match the database")

        after = self._snapshot()
        if before != after:
            raise CommandError(f"Excel export mutated financial data: before={before} after={after}")

        self.stdout.write("EXPENSE EXPORT V6: XLSX opens successfully")
        self.stdout.write("EXPENSE EXPORT V6: five financial sheets are present")
        self.stdout.write("EXPENSE EXPORT V6: expense/receivable row counts match DB")
        self.stdout.write("EXPENSE EXPORT V6: export is read-only")
        self.stdout.write(self.style.SUCCESS("SUCCESS: EXPENSE TRACKER FINANCIAL EXPORT V6 REGRESSION PASSED"))
